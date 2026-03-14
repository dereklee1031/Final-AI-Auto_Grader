from __future__ import annotations

from pathlib import Path
from typing import Any

from core.chunking import clean_text, make_sort_key
from extractors.html_extractor import (
    ExtractedHTML,
    HtmlTextBlock,
    extracted_html_to_jsonable,
)


def extract_pptx(file_path: Path, rel_path: str, cfg: dict[str, Any]) -> ExtractedHTML:
    from pptx import Presentation  # type: ignore

    prs = Presentation(str(file_path))
    min_chars = int(cfg.get("min_text_chars", 30))
    out: list[HtmlTextBlock] = []
    stats = {"slides_scanned": 0, "shapes_scanned": 0, "text_blocks": 0}
    doc_order = 0
    block_idx = 0

    for slide_idx, slide in enumerate(prs.slides, start=1):
        stats["slides_scanned"] += 1
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            stats["shapes_scanned"] += 1
            for para in shape.text_frame.paragraphs:
                text = clean_text("".join(run.text or "" for run in para.runs))
                if len(text) >= min_chars:
                    block_idx += 1
                    doc_order += 1
                    out.append(
                        HtmlTextBlock(
                            element_tag="p",
                            page_number=slide_idx,
                            block_index=block_idx,
                            text=text,
                            sort_key=make_sort_key(slide_idx, block_idx),
                            document_order=doc_order,
                        )
                    )

    stats["text_blocks"] = len(out)
    return ExtractedHTML(source_path=rel_path, file_type="pptx", text_blocks=out, stats=stats)


extracted_pptx_to_jsonable = extracted_html_to_jsonable
